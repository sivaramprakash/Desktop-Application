import sys
import os
from docx import Document
from openai import OpenAI
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

def print_help():
    print("Usage: python generate_doc.py \"<template_path>\" \"<topic>\"")
    print("Example: python generate_doc.py \"C:\\Users\\Documents\\template.docx\" \"Artificial Intelligence in Healthcare\"")

def initialize_ai():
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key or api_key.strip() == "your_openai_api_key_here":
        print("Error: OPENAI_API_KEY not found or invalid. Please add it to your .env file.")
        sys.exit(1)
    return OpenAI(api_key=api_key)

def generate_new_text(client: OpenAI, original_text: str, topic: str) -> str:
    """
    Uses AI to generate completely new textual content based on the original structure, length, and a new topic.
    """
    if not original_text.strip():
        return ""
        
    system_prompt = (
        "You are an expert ghostwriter. Your task is to rewrite the provided text "
        f"so that it is entirely about the topic [{topic}].\n"
        "CRITICAL RULES:\n"
        "1. DO NOT change the structural nature (e.g., if it's a short title, output a short title. If it's a 3-sentence paragraph, output a 3-sentence paragraph).\n"
        "2. Keep the word count roughly identical.\n"
        "3. Output ONLY the new text, without any conversational filler, markdown formatting or quotes.\n"
        "4. Do NOT say 'Here is the rewritten text:'."
    )
    
    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": original_text}
            ],
            temperature=0.4,
            timeout=3.0
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        error_msg = str(e).lower()
        if "quota" in error_msg or "429" in error_msg or "timeout" in error_msg:
             print(f"Warning: API Issue (Quota/Timeout). Replacing '{original_text[:15]}...' with dummy text for {topic}.")
             # Mock replacement to demonstrate that the format preservation still works
             words_len = len(original_text.split())
             return f"[AI Content for {topic}] " * (words_len // 4 + 1)
             
        print(f"Error generating text for '{original_text[:20]}...': {e}")
        return original_text # Fallback to original if AI fails

def replace_paragraph_text(paragraph, new_text):
    """
    Replaces the text of a paragraph while attempting to preserve its original font styling, alignment, and spacing.
    We place all the new text into the first Run object and clear the rest, preserving the base paragraph and first-run style.
    """
    if not paragraph.runs or not new_text:
        return
        
    # Put all new text in the first run to preserve the primary run styling (font, bold, italic, size)
    paragraph.runs[0].text = new_text
    
    # Clear remaining runs to remove the old text, but keep the runs themselves if they have complex xml structure underneath
    for i in range(1, len(paragraph.runs)):
        paragraph.runs[i].text = ""

def process_document(template_path: str, topic: str):
    if not os.path.exists(template_path):
        print(f"Error: The template file '{template_path}' does not exist.")
        sys.exit(1)
        
    if not template_path.lower().endswith(".docx"):
        print("Error: The template file must be a Microsoft Word .docx file.")
        sys.exit(1)
        
    print(f"-> Loading template: {template_path}")
    try:
        doc = Document(template_path)
    except Exception as e:
        print(f"Error loading document: {e}")
        sys.exit(1)
        
    client = initialize_ai()
    
    print(f"-> Generating new content about '{topic}'...")
    
    # Process standard paragraphs
    print("   -> Processing paragraphs...")
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            # Generate new text
            new_text = generate_new_text(client, text, topic)
            # Replace while preserving style
            replace_paragraph_text(para, new_text)
            
    # Process tables (often used for layouts, headers, or strict formatted structures)
    print("   -> Processing tables...")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    text = para.text.strip()
                    if text:
                        new_text = generate_new_text(client, text, topic)
                        replace_paragraph_text(para, new_text)
                        
    # Sections (Headers and Footers)
    print("   -> Processing headers and footers...")
    for section in doc.sections:
        # Header
        for para in section.header.paragraphs:
            text = para.text.strip()
            if text:
                new_text = generate_new_text(client, text, topic)
                replace_paragraph_text(para, new_text)
        # Footer
        for para in section.footer.paragraphs:
            text = para.text.strip()
            if text:
                new_text = generate_new_text(client, text, topic)
                replace_paragraph_text(para, new_text)

    # Save output to Downloads folder
    safe_topic = "".join([c if c.isalnum() else "_" for c in topic])
    output_filename = f"generated_{safe_topic}.docx"
    
    # Get the path to the user's Downloads directory
    downloads_dir = os.path.join(os.path.expanduser("~"), "Downloads")
    output_path = os.path.join(downloads_dir, output_filename)
    
    try:
        doc.save(output_path)
        print(f"\n-> Success! Generated document saved to: {output_path}")
    except Exception as e:
        print(f"\n-> Error saving document to {output_path}: {e}")

if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Error: Invalid number of arguments.")
        print_help()
        sys.exit(1)
        
    template_path = sys.argv[1]
    topic = sys.argv[2]
    
    process_document(template_path, topic)
