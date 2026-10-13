"""
ppt_generator.py
Standalone script to generate a PowerPoint presentation using python-pptx.
Usage: python ppt_generator.py "<topic>" <num_slides> "<save_path>"
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def add_title_slide(prs: Presentation, topic: str):
    """Creates the first title slide."""
    layout = prs.slide_layouts[0]  # Title Slide layout
    slide = prs.slides.add_slide(layout)

    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = topic
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 56, 168)

    subtitle.text = f"A Presentation on {topic}"
    subtitle.text_frame.paragraphs[0].font.size = Pt(20)
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(64, 64, 64)


def add_content_slide(prs: Presentation, heading: str, bullets: list):
    """Creates a content slide with a heading and bullet points."""
    layout = prs.slide_layouts[1]  # Title and Content layout
    slide = prs.slides.add_slide(layout)

    title = slide.shapes.title
    body = slide.placeholders[1]

    title.text = heading
    title.text_frame.paragraphs[0].font.size = Pt(32)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 56, 168)

    tf = body.text_frame
    tf.word_wrap = True

    for i, bullet in enumerate(bullets):
        if i == 0:
            tf.text = bullet
            tf.paragraphs[0].font.size = Pt(18)
        else:
            p = tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(18)
            p.level = 0


def build_slide_content(topic: str, num_slides: int) -> list:
    """
    Builds a list of (heading, [bullets]) tuples for num_slides content slides.
    This is a template-based fallback that does not require an API.
    """
    # Generic slide templates adaptable to any topic
    templates = [
        ("Introduction", [
            f"Overview of {topic}",
            "Background and context",
            "Why this topic matters today",
        ]),
        ("Key Concepts", [
            f"Core principles of {topic}",
            "Foundational terminology",
            "Essential ideas to understand",
        ]),
        ("Current Trends", [
            f"Latest developments in {topic}",
            "Industry adoption and growth",
            "Emerging innovations",
        ]),
        ("Applications & Use Cases", [
            f"Real-world applications of {topic}",
            "Industry-specific implementations",
            "Success stories and examples",
        ]),
        ("Challenges & Limitations", [
            f"Common obstacles in {topic}",
            "Technical or practical limitations",
            "Areas requiring further research",
        ]),
        ("Future Outlook", [
            f"The road ahead for {topic}",
            "Predicted advancements",
            "Opportunities and potential impact",
        ]),
        ("Summary & Conclusion", [
            f"Key takeaways about {topic}",
            "Recap of important points",
            "Final thoughts and recommendations",
        ]),
    ]

    # Pick as many templates as needed (cycle if num_slides > len(templates))
    slides = []
    for i in range(num_slides):
        slides.append(templates[i % len(templates)])
    return slides


def generate_ppt(topic: str, num_slides: int, save_path: str) -> str:
    """Main function to generate the PowerPoint file."""
    prs = Presentation()

    # Set slide dimensions to widescreen (16:9)
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title
    add_title_slide(prs, topic)

    # Remaining slides: Content
    content_slides = build_slide_content(topic, num_slides)
    for heading, bullets in content_slides:
        add_content_slide(prs, heading, bullets)

    # Ensure save directory exists
    os.makedirs(os.path.dirname(os.path.abspath(save_path)), exist_ok=True)

    prs.save(save_path)
    print(f"Presentation saved to: {save_path}")
    return save_path


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python ppt_generator.py \"<topic>\" <num_slides> \"<save_path>\"")
        sys.exit(1)

    topic_arg = sys.argv[1]
    num_slides_arg = int(sys.argv[2])
    save_path_arg = sys.argv[3]

    output = generate_ppt(topic_arg, num_slides_arg, save_path_arg)

    # Auto-open the file
    os.startfile(output)
    print("Done.")
